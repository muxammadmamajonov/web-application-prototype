"""SEC-010 — CSV/XLSX exports must not hand a spreadsheet an executable formula.

`?format=csv|xlsx` is an active MVP export served as an attachment. A cell whose
first effective character is `=`, `+`, `-` or `@` is interpreted as a formula by
Excel / LibreOffice / Sheets when the file is opened, so a value typed into the
CRM (a customer name, a note) becomes code on the operator's machine.

One shared predicate neutralises both writers; numerics stay numeric.
"""
from __future__ import annotations

import csv
import io
from datetime import date
from decimal import Decimal
from uuid import UUID

import pytest

from app.api.reports import _csv_bytes, _report_bytes, _safe_cell, _xlsx_bytes

DANGEROUS = ["=1+1", "+1", "-1+2", "@SUM(A1)", "=cmd|'/c calc'!A0"]
WHITESPACE_LED = ["   =1+1", "  @SUM(A1)", " \t=1+1"]
CONTROL_LED = ["\t=1+1", "\r=1+1", "\n=1+1", "\tplain"]
SAFE_STRINGS = ["Alice", "Ali =bek", "ALI", "e=mc2", "", "Алишер", "客户", "+998 90 123 45 67".replace("+", "")]


# ── the shared predicate ────────────────────────────────────────────────────
@pytest.mark.parametrize("value", DANGEROUS)
def test_dangerous_leading_marker_is_prefixed(value: str) -> None:
    out = _safe_cell(value)
    assert out == "'" + value


@pytest.mark.parametrize("value", WHITESPACE_LED)
def test_leading_whitespace_does_not_bypass(value: str) -> None:
    """`   =1+1` still evaluates in a spreadsheet, so it must be neutralised."""
    out = _safe_cell(value)
    assert out.startswith("'"), f"whitespace-led payload bypassed: {value!r}"
    assert out == "'" + value


@pytest.mark.parametrize("value", CONTROL_LED)
def test_control_character_lead_is_neutralised(value: str) -> None:
    out = _safe_cell(value)
    assert out == "'" + value


@pytest.mark.parametrize("value", SAFE_STRINGS)
def test_safe_strings_are_untouched(value: str) -> None:
    assert _safe_cell(value) == value


def test_neutralisation_is_idempotent() -> None:
    """Re-running the predicate must not stack apostrophes."""
    once = _safe_cell("=1+1")
    twice = _safe_cell(once)
    assert once == "'=1+1"
    assert twice == once


def test_non_string_scalars_keep_their_type() -> None:
    """Numerics must stay numeric — including negatives, which start with '-'."""
    for value in (0, 5, -5, 1234567, Decimal("-12.34"), Decimal("0.00"), 3.5, -3.5, True, False, None):
        assert _safe_cell(value) is value, f"{value!r} was coerced"


def test_dates_and_uuids_are_not_stringified() -> None:
    d = date(2026, 7, 29)
    u = UUID("0c1d2e3f-aaaa-bbbb-cccc-000011112222")
    assert _safe_cell(d) is d
    assert _safe_cell(u) is u


# ── CSV writer ──────────────────────────────────────────────────────────────
COLUMNS = [("name", "Name"), ("amount", "Amount")]


def test_csv_output_parses_and_carries_a_literal_apostrophe() -> None:
    rows = [{"name": "=cmd|'/c calc'!A0", "amount": -5}]
    text = _csv_bytes(COLUMNS, rows).decode("utf-8-sig")
    parsed = list(csv.reader(io.StringIO(text)))
    assert parsed[0] == ["Name", "Amount"], "static headers must not change"
    cell = parsed[1][0]
    assert cell.startswith("'"), f"csv cell not neutralised: {cell!r}"
    assert cell == "'=cmd|'/c calc'!A0"
    assert parsed[1][1] == "-5", "negative numeric must stay numeric in csv"


def test_csv_safe_rows_are_unchanged() -> None:
    rows = [{"name": "Alice", "amount": 100}]
    parsed = list(csv.reader(io.StringIO(_csv_bytes(COLUMNS, rows).decode("utf-8-sig"))))
    assert parsed[1] == ["Alice", "100"]


# ── XLSX writer ─────────────────────────────────────────────────────────────
def _load(data: bytes):
    from openpyxl import load_workbook
    return load_workbook(io.BytesIO(data)).active


def test_xlsx_dangerous_cell_is_an_inert_string() -> None:
    rows = [{"name": "=1+1", "amount": -5}]
    ws = _load(_xlsx_bytes("Report", COLUMNS, rows))
    assert [c.value for c in ws[1]] == ["Name", "Amount"], "static headers must not change"
    cell = ws.cell(row=2, column=1)
    assert cell.data_type != "f", "cell is still a formula"
    assert isinstance(cell.value, str)
    assert cell.value == "'=1+1"


def test_xlsx_numeric_stays_numeric() -> None:
    ws = _load(_xlsx_bytes("Report", COLUMNS, [{"name": "Alice", "amount": -5}]))
    amount = ws.cell(row=2, column=2)
    assert amount.data_type == "n"
    assert amount.value == -5


# ── the real export boundary ────────────────────────────────────────────────
def test_report_bytes_csv_neutralises_at_the_service_boundary() -> None:
    data, fname, media = _report_bytes("csv", "aging", "Aging", COLUMNS, [{"name": "@SUM(A1)", "amount": 1}])
    assert media == "text/csv" and fname.endswith(".csv")
    parsed = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    assert parsed[1][0] == "'@SUM(A1)"


def test_report_bytes_xlsx_neutralises_at_the_service_boundary() -> None:
    data, fname, media = _report_bytes("xlsx", "aging", "Aging", COLUMNS, [{"name": "+1", "amount": 1}])
    assert fname.endswith(".xlsx")
    assert media.endswith("spreadsheetml.sheet")
    cell = _load(data).cell(row=2, column=1)
    assert cell.data_type != "f"
    assert cell.value == "'+1"
